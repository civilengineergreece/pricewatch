# -*- coding: utf-8 -*-
"""
PriceWatch - ωριαία παρακολούθηση τιμών (Plaisio) με ειδοποίηση στο Telegram.
Δωρεάν. Τρέχει σε Windows (.exe) ή με Python.

Χρήση:
  PriceWatch.exe            -> τρέχει συνέχεια, έλεγχος κάθε ώρα
  PriceWatch.exe --once     -> ένας έλεγχος και τέλος (π.χ. για Task Scheduler)
  PriceWatch.exe --test     -> δοκιμή: διαβάζει τη σελίδα, δείχνει τι βρήκε, ΔΕΝ στέλνει τίποτα
"""
import base64
import csv
import html
import json
import os
import random
import re
import string
import sqlite3
import statistics
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from urllib.parse import quote, urljoin, urlparse, parse_qsl, urlencode, urlunparse

import requests
from bs4 import BeautifulSoup

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

APP_DIR = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))
DATA_DIR = os.environ.get("PRICEWATCH_DATA_DIR") or APP_DIR
os.makedirs(DATA_DIR, exist_ok=True)
CONFIG_PATH = os.path.join(APP_DIR, "config.json")
DB_PATH = os.path.join(DATA_DIR, "prices.db")
LOG_PATH = os.path.join(DATA_DIR, "pricewatch.log")

DEFAULT_CONFIG = {
    "telegram_token": "",
    "telegram_chat_id": "",
    "start_url": "https://www.plaisio.gr/list/obtainability--online?pagesize=96&sorting=saveAmount_desc",
    "page_url_template": "https://www.plaisio.gr/list/obtainability--online?pagesize=96&sorting=saveAmount_desc&page={n}",
    "max_pages": 105,
    "page_delay_sec": 2.0,
    "check_every_minutes": 60,
    "min_discount_percent": 50,
    "min_drop_vs_history_percent": 50,
    "min_new_price_eur": 20,
    "max_alerts_per_run": 15,
    "report_min_price_eur": 5,
    "daily_digest_top_n": 15,
    "public_enabled": True,
    "public_api_url": "https://app.findbar.io/api/v1/search?locale=el",
    "public_site": "https://www.public.gr",
    "public_page_size": 90,
    "public_max_pages": 8,
    "public_stop_below_percent": 30,
    "public_bearer": "",
    "daily_xlsx_to_telegram": False,
    "public_exclude_categories": ["Μουσική", "Βιβλί"],
    "kotsovolos_enabled": True,
    "kotsovolos_site": "https://www.kotsovolos.gr",
    "kotsovolos_page_size": 48,
    "kotsovolos_max_pages": 40,
    "kotsovolos_catalog_id": "10551",
    "kotsovolos_lang_id": "-24",
    "kotsovolos_store_id": "10151",
    "kotsovolos_order_by": "10",
    "kotsovolos_categories": ["35261"],
    "kotsovolos_category_paths": [
        "/computing/laptop-tablet-ipad/notebook-macbook-ultrabook",
        "/computing/laptop-tablet-ipad/tablets-ipad",
        "/computing/monitors",
        "/computing/desktops-and-upgrade-tools/desktops-mac-all-in-one",
        "/computing/printing-consumables/printers",
        "/computing/storage-hard-disk-drives-and-usb-sticks/external-ssd",
        "/sound-vision/televisions/led-lcd",
        "/sound-vision/home-cinema/soundbars",
        "/mobile-phones-gps/wearables/smartwatches",
        "/mobile-phones-gps/akoystika/akoustika-earbuds",
        "/mobile-phones-gps/akoystika/akoustika-kefalis",
        "/gaming-gadgets/playstation-5/konsola-ps5",
        "/gaming-gadgets/nintendo-switch-2/consoles-nintendo-switch-2",
        "/imaging/digital-cameras/fotografikes-mirrorless",
        "/imaging/action/cameras",
        "/gaming-gadgets/thlekatefthinomena/drones",
    ],
    "germanos_enabled": True,
    "germanos_max_pages": 30,
    "germanos_sort": "&sort=defaultPrice_ascending&sortType=numprop",
    "germanos_urls": [
        "https://www.germanos.gr/category/telephony/mobile/smartphones-all/?categoryId=cat310005",
        "https://www.germanos.gr/category/laptops-peripherals/laptops/?categoryId=cat310150",
        "https://www.germanos.gr/category/televisions/televisions/?categoryId=cat310016",
        "https://www.germanos.gr/category/telephony/all-tablets/?categoryId=cat310144",
        "https://www.germanos.gr/category/wearables-gadgets/wearables/smartwatches/?categoryId=cat310078",
        "https://www.germanos.gr/category/gaming/consoles/?categoryId=cat310062",
        "https://www.germanos.gr/category/laptops-peripherals/computers-peripherals/monitors/?categoryId=cat1050038",
    ],
}

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml",
    "Accept-Language": "el-GR,el;q=0.9,en;q=0.8",
}

CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\u200b-\u200f\u2028\u2029\ufeff]")


def clean_str(v):
    """Αφαιρεί αόρατους/ελεγκτικούς χαρακτήρες που δεν δέχεται το Excel και καθαρίζει τα κενά."""
    return re.sub(r"\s+", " ", CTRL_RE.sub(" ", v)).strip()


CODE_RE = re.compile(r"Κωδ\.\s*Πλαίσιο\s*(\d+)")
CARD_RE = re.compile(
    r"([\d.]+,\d{2})\s*€\s*([\d.]+,\d{2})\s*€\s*Κέρδος\s*([\d.]+,\d{2})\s*€"
    r"(?:\s*Προσθήκη)?\s*(.*?)\s*Κωδ\.\s*Πλαίσιο\s*(\d+)",
    re.S,
)


PLAISIO_SLUGS = {
    "pc-perifereiaka": "Υπολογιστές & περιφερειακά", "laptops": "Laptops", "othones": "Οθόνες",
    "oles-oi-othones": "Όλες οι οθόνες", "desktops": "Desktops", "all-in-one-pcs": "All-in-One PCs",
    "anavathmisi-diktia": "Αναβάθμιση & δίκτυα", "diktiaka": "Δικτυακά", "smarthome": "Smart home",
    "voip-conference": "VoIP & συνέδρια", "anavathmisi-pc": "Αναβάθμιση PC",
    "anemistirakia-psiktres": "Ανεμιστηράκια & ψύκτρες", "koutia-desktop": "Κουτιά desktop",
    "ilektrologika-ergaleia": "Ηλεκτρολογικά & εργαλεία", "categories-ergaleia": "Εργαλεία", "ups": "UPS",
    "ektiposi-analwsima": "Εκτύπωση & αναλώσιμα", "ektipotika-mixanimata": "Εκτυπωτικά μηχανήματα",
    "polimixanimata": "Πολυμηχανήματα", "megales-oikiakes-siskeves": "Μεγάλες οικιακές συσκευές",
    "plisi-stegnoma-rouxon": "Πλύση & στέγνωμα ρούχων", "plintiria-rouxon": "Πλυντήρια ρούχων",
    "psigeia-sintirisi": "Ψυγεία & συντήρηση", "diporta-psigeia": "Δίπορτα ψυγεία",
    "mikres-oikiakes-siskeves": "Μικρές οικιακές συσκευές", "kathariotita": "Καθαριότητα",
    "skoupes-rompot": "Σκούπες ρομπότ", "thleoraseis-ixos": "Τηλεοράσεις & ήχος", "thleoraseis": "Τηλεοράσεις",
    "tilefonia-tablet-wearables": "Τηλεφωνία, tablet & wearables", "tilefona": "Τηλέφωνα",
    "smartphones": "Smartphones", "prosopiki-peripoiisi": "Προσωπική περιποίηση",
    "peripoiisi-prosopou": "Περιποίηση προσώπου", "multigrooming-set": "Multigrooming set",
    "fotografia-drones-ilektrokinisi": "Φωτογραφία, drones & ηλεκτροκίνηση", "fotografia": "Φωτογραφία",
    "aksesouar-fotografikon-mixanimon": "Αξεσουάρ φωτογραφικών", "klima-spitiou": "Κλίμα & σπίτι",
    "afigrantires-katharistes-aera": "Αφυγραντήρες & καθαριστές αέρα",
    "ionistes-katharistes-aera": "Ιονιστές & καθαριστές αέρα",
}


def pretty_slug(seg):
    return PLAISIO_SLUGS.get(seg) or seg.replace("-", " ").replace("_", " ").strip().capitalize()


def plaisio_category(url):
    m = re.search(r"/product/([^?#]+)", url or "")
    if not m:
        return ""
    segs = [x for x in m.group(1).split("/") if x][:-1][:3]
    return " › ".join(pretty_slug(x) for x in segs)


def log(msg):
    line = f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {msg}"
    print(line, flush=True)
    try:
        with open(LOG_PATH, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass


# ----------------------------------------------------------------- config
def load_config():
    if not os.path.exists(CONFIG_PATH):
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(DEFAULT_CONFIG, f, ensure_ascii=False, indent=2)
        log(f"Δημιουργήθηκε το {CONFIG_PATH}. Βάλε μέσα το telegram_token και ξανατρέξε.")
    with open(CONFIG_PATH, encoding="utf-8") as f:
        cfg = json.load(f)
    missing = [k for k in DEFAULT_CONFIG if k not in cfg]
    for k, v in DEFAULT_CONFIG.items():
        cfg.setdefault(k, v)
    if missing:
        save_config(cfg)
    # μεταβλητές περιβάλλοντος (cloud): δεν γράφονται ποτέ στο config.json
    for env, key in (("PW_TELEGRAM_TOKEN", "telegram_token"), ("PW_TELEGRAM_CHAT_ID", "telegram_chat_id"),
                     ("PW_PUBLIC_BEARER", "public_bearer")):
        if os.environ.get(env, "").strip():
            cfg[key] = os.environ[env].strip()
    if os.environ.get("PW_DAILY_XLSX"):
        cfg["daily_xlsx_to_telegram"] = True
    return cfg


def save_config(cfg):
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


# --------------------------------------------------------------- telegram
def tg_send(cfg, text):
    if not cfg["telegram_token"] or not cfg["telegram_chat_id"]:
        log("Δεν έχει ρυθμιστεί Telegram - το μήνυμα δεν στάλθηκε.")
        return False
    try:
        r = requests.post(
            f"https://api.telegram.org/bot{cfg['telegram_token']}/sendMessage",
            data={"chat_id": cfg["telegram_chat_id"], "text": text, "parse_mode": "HTML",
                  "disable_web_page_preview": "true"},
            timeout=30,
        )
        if r.status_code != 200:
            log(f"Telegram σφάλμα {r.status_code}: {r.text[:200]}")
            return False
        return True
    except Exception as e:
        log(f"Telegram σφάλμα: {e}")
        return False


def tg_send_document(cfg, path, caption=""):
    if not cfg["telegram_token"] or not cfg["telegram_chat_id"] or not os.path.exists(path):
        return False
    try:
        with open(path, "rb") as f:
            r = requests.post(f"https://api.telegram.org/bot{cfg['telegram_token']}/sendDocument",
                              data={"chat_id": cfg["telegram_chat_id"], "caption": caption},
                              files={"document": (os.path.basename(path), f)}, timeout=120)
        if r.status_code != 200:
            log(f"Telegram (αρχείο) σφάλμα {r.status_code}")
            return False
        return True
    except Exception as e:
        log("Telegram (αρχείο) σφάλμα: " + str(e).replace(cfg["telegram_token"], "***"))
        return False


def detect_chat_id(cfg):
    """Βρίσκει μόνο του το chat id: στείλε οποιοδήποτε μήνυμα στο bot σου και τρέξε το πρόγραμμα."""
    if cfg["telegram_chat_id"] or not cfg["telegram_token"]:
        return
    try:
        r = requests.get(f"https://api.telegram.org/bot{cfg['telegram_token']}/getUpdates", timeout=30)
        for upd in reversed(r.json().get("result", [])):
            msg = upd.get("message") or upd.get("channel_post")
            if msg:
                cfg["telegram_chat_id"] = str(msg["chat"]["id"])
                save_config(cfg)
                log("Βρέθηκε το Telegram chat id και αποθηκεύτηκε.")
                return
        log("Δεν βρέθηκε μήνυμα. Στείλε ένα μήνυμα (π.χ. 'hi') στο bot σου στο Telegram και ξανατρέξε.")
    except Exception as e:
        log(f"Αδυναμία εύρεσης chat id: {e}")


# ----------------------------------------------------------------- parsing
def to_float(s):
    return float(s.replace(".", "").replace(",", "."))


def parse_products(page_html, page_url):
    """Διαβάζει τις κάρτες προϊόντων. Δεν βασίζεται σε CSS classes, μόνο στο κείμενο
    'Κωδ. Πλαίσιο NNNNNNN' κάθε κάρτας, ώστε να αντέχει σε αλλαγές του site."""
    soup = BeautifulSoup(page_html, "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()

    anchors = {}
    for el in soup.find_all(True):
        txt = el.get_text(" ", strip=True)
        if len(txt) < 60:
            m = CODE_RE.search(txt)
            if m and m.group(1) not in anchors:
                anchors[m.group(1)] = el

    products = {}
    for code, el in anchors.items():
        card = el
        while card.parent is not None and len(CODE_RE.findall(card.parent.get_text(" ", strip=True))) <= 1:
            card = card.parent
        text = clean_str(card.get_text(" ", strip=True))
        m = CARD_RE.search(text)
        if not m:
            continue  # προϊόν χωρίς έκπτωση / άγνωστη μορφή
        old, new, saving, name, code2 = to_float(m.group(1)), to_float(m.group(2)), to_float(m.group(3)), \
            m.group(4).strip(), m.group(5)
        if code2 != code or new <= 0 or old <= new:
            continue
        link = page_url
        for a in card.find_all("a", href=True):
            h = a["href"]
            if h.startswith(("#", "javascript", "mailto")) or re.search(r"wishlist|compare|cart", h, re.I):
                continue
            link = urljoin(page_url, h)
            if name and name[:12].lower() in a.get_text(" ", strip=True).lower():
                break
        products[code] = {
            "code": code, "name": name[:140], "old": old, "new": new, "saving": saving,
            "url": link, "clearance": "CLEARANCE" in text, "store": "Plaisio", "category": plaisio_category(link),
        }
    return products


def find_page_url_builder(cfg, page_html, base_url):
    """Βρίσκει πώς χτίζεται το URL της σελίδας N (από config ή από τα links της σελιδοποίησης)."""
    tpl = cfg.get("page_url_template", "").strip()
    if tpl:
        return lambda n: tpl.replace("{n}", str(n))
    soup = BeautifulSoup(page_html, "html.parser")
    for a in soup.find_all("a", href=True):
        if a.get_text(strip=True) == "2":
            href = urljoin(base_url, a["href"])
            p = urlparse(href)
            q = parse_qsl(p.query, keep_blank_values=True)
            for k, v in q:
                if v == "2":
                    def build(n, p=p, q=q, k=k):
                        qq = [(kk, str(n) if kk == k else vv) for kk, vv in q]
                        return urlunparse(p._replace(query=urlencode(qq)))
                    return build
            if re.search(r"(?<!\d)2(?!\d)", p.path):
                return lambda n, p=p: urlunparse(p._replace(path=re.sub(r"(?<!\d)2(?!\d)", str(n), p.path, count=1)))
    return None


def fetch(url):
    r = requests.get(url, headers=HEADERS, timeout=45)
    if r.status_code != 200:
        raise RuntimeError(f"HTTP {r.status_code} για {url}")
    return r.text


def scan(cfg, test=False):
    start = cfg["start_url"]
    first_html = fetch(start)
    if test:
        with open(os.path.join(APP_DIR, "page1_debug.html"), "w", encoding="utf-8") as f:
            f.write(first_html)
    all_products = parse_products(first_html, start)
    log(f"Plaisio σελίδα 1: {len(all_products)} προϊόντα με έκπτωση")
    if not all_products:
        raise RuntimeError("Δεν διάβασα κανένα προϊόν. Η σελίδα άλλαξε μορφή ή με μπλόκαρε το site.")

    build = find_page_url_builder(cfg, first_html, start)
    if build is None:
        log("ΠΡΟΣΟΧΗ: δεν βρήκα πώς αλλάζει η σελίδα (page 2, 3...). Ελέγχω μόνο τη σελίδα 1. "
            "Βάλε στο config.json το page_url_template με {n}, π.χ. ...&pageno={n}")
    else:
        pages = 1 if test else cfg["max_pages"]
        for n in range(2, (3 if test else pages) + 1):
            time.sleep(cfg["page_delay_sec"])
            try:
                ph = fetch(build(n))
            except Exception as e:
                log(f"Σελίδα {n}: {e} - σταματώ τη σάρωση εδώ")
                break
            found = parse_products(ph, build(n))
            new_codes = [c for c in found if c not in all_products]
            if not new_codes:
                log(f"Plaisio σελίδα {n}: τίποτα καινούργιο - τέλος σάρωσης")
                break
            all_products.update({c: found[c] for c in new_codes})
    return all_products


# ============================================================== PUBLIC
PUBLIC_NEW = ["finalprice", "offerprice", "saleprice", "specialprice", "discountedprice", "currentprice",
              "sellingprice", "pricewithvat", "price"]
PUBLIC_OLD = ["originalprice", "oldprice", "listprice", "regularprice", "initialprice", "retailprice",
              "baseprice", "fullprice", "wasprice", "pricebeforediscount", "msrp", "rrp", "recommendedprice"]
PUBLIC_PCT = ["autodiscountpercent", "discountpercent", "discountpercentage", "discountpct", "percentoff",
              "percentdiscount", "discountrate"]
PUBLIC_NAME = ["title", "name", "productname", "displayname"]
PUBLIC_URL = ["url", "producturl", "seourl", "canonicalurl", "link", "permalink", "href", "path", "slug"]
PUBLIC_CODE = ["sku", "uniqueid", "partnumber", "articlenumber", "id", "productid", "code", "itemid", "skuid"]


def _nk(k):
    return re.sub(r"[^a-z0-9]", "", str(k).lower())


def _price_list(v):
    """Λίστα τιμών τύπου [{"usage":"Display","value":...},{"usage":"Offer","value":...}] -> (offer, display, όλες)."""
    offer = display = None
    vals = []
    for x in v:
        if not isinstance(x, dict):
            continue
        val = _num(x.get("value", x.get("price")))
        if val is None:
            continue
        vals.append(val)
        u = str(x.get("usage", x.get("type", ""))).lower()
        if u == "offer":
            offer = val
        elif u == "display":
            display = val
    return offer, display, vals


def _price_max(v):
    if isinstance(v, list):
        _, _, vals = _price_list(v)
        return max(vals) if vals else None
    return _num(v)


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, list):
        offer, display, vals = _price_list(v)
        if offer is not None:
            return offer
        return vals[0] if vals else None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, dict):
        for k in ("amount", "value", "price", "raw", "final"):
            if k in v:
                r = _num(v[k])
                if r is not None:
                    return r
        return None
    if isinstance(v, str):
        t = re.sub(r"[^\d.,-]", "", v)
        if not t:
            return None
        if "," in t and "." in t:
            t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
        elif "," in t:
            t = t.replace(",", ".")
        try:
            return float(t)
        except ValueError:
            return None
    return None


def _text(v):
    if isinstance(v, str) and clean_str(v):
        return clean_str(v)
    if isinstance(v, dict):
        for k in ("el", "name", "title", "value"):
            if isinstance(v.get(k), str) and v[k].strip():
                return v[k].strip()
    if isinstance(v, (int,)) and not isinstance(v, bool):
        return str(v)
    return None


def _find(obj, names, conv, depth=3):
    wanted = {n: i for i, n in enumerate(names)}
    queue, best = [(obj, 0)], None
    while queue:
        o, d = queue.pop(0)
        if isinstance(o, dict):
            for k, v in o.items():
                kl = _nk(k)
                if kl in wanted:
                    r = conv(v)
                    if r is not None and (best is None or wanted[kl] < best[0]):
                        best = (wanted[kl], r)
                if isinstance(v, (dict, list)) and d < depth:
                    queue.append((v, d + 1))
        elif isinstance(o, list) and d < depth:
            for v in o[:5]:
                if isinstance(v, (dict, list)):
                    queue.append((v, d + 1))
    return None if best is None else best[1]


def _src(it):
    """Κάθε προϊόν έρχεται ως {"id":..., "source": {...πεδία...}} - παίρνουμε τα πεδία του."""
    if isinstance(it, dict) and isinstance(it.get("source"), dict):
        d = dict(it["source"])
        d.setdefault("_hit_id", it.get("id"))
        return d
    return it


def _public_name(d, depth=1):
    n = _find(d, PUBLIC_NAME, _text, depth=depth)
    if n:
        return n
    for k, v in d.items():
        kl = _nk(k)
        if ("title" in kl or kl.endswith("name")) and not any(x in kl for x in ("brand", "author", "categor", "manufact", "seller")):
            t = _text(v)
            if t:
                return t
    return None


def _looks_like_product(it, depth=1):
    d = _src(it)
    return isinstance(d, dict) and _public_name(d, depth) is not None and _find(d, PUBLIC_NEW, _num, depth=depth) is not None


def _find_items(data, depth=1):
    best = []
    stack = [data]
    while stack:
        o = stack.pop()
        if isinstance(o, list):
            dicts = [x for x in o if isinstance(x, dict)]
            if dicts and _looks_like_product(dicts[0], depth) and len(dicts) > len(best):
                best = dicts
            for x in o[:3]:
                if isinstance(x, (dict, list)):
                    stack.append(x)
        elif isinstance(o, dict):
            for v in o.values():
                if isinstance(v, (dict, list)):
                    stack.append(v)
    return best


def _public_excluded(d, cfg):
    words = [w.lower() for w in cfg.get("public_exclude_categories", []) if w]
    if not words:
        return False
    crumbs = d.get("breadcrumbs")
    names = []
    if isinstance(crumbs, list):
        for c in crumbs:
            if isinstance(c, dict):
                names.append(str(c.get("displayName", "")).lower())
    return any(w in n for w in words for n in names)


def parse_public(data, site, cfg=None, store="Public", prefix="public:", depth=1, fallback=None):
    cfg = cfg or DEFAULT_CONFIG
    products = {}
    for hit in _find_items(data, depth):
        it = _src(hit)
        if _public_excluded(it, cfg):
            continue
        name = _public_name(it, depth)
        new = _find(it, PUBLIC_NEW, _num, depth=depth)
        old = _find(it, PUBLIC_OLD, _num, depth=depth)
        pct = _find(it, PUBLIC_PCT, _num, depth=depth)
        if not name or not new or new <= 0:
            continue
        if old is None:
            base = _find(it, ["price"], _price_max, depth=depth)
            if base and base > new:
                old = base
        if old is not None and old < new:
            old, new = new, old
        if (old is None or old <= new) and pct and 0 < pct < 100:
            old = new / (1 - pct / 100)
        if old is None or old <= new:
            continue
        code = _find(it, PUBLIC_CODE, _text, depth=depth) or str(hit.get("id") or abs(hash(name)))
        link = _find(it, PUBLIC_URL, _text, depth=depth) or ""
        if link:
            link = urljoin(site + "/", link)
        else:
            link = fallback or f"{site}/search?q={quote(name[:60])}"
        crumbs = it.get("breadcrumbs")
        cat = ""
        if isinstance(crumbs, list):
            cat = " › ".join(clean_str(str(c.get("displayName", ""))) for c in crumbs[:3] if isinstance(c, dict))
        products[prefix + code] = {
            "code": prefix + code, "name": name[:140], "old": round(old, 2), "new": round(new, 2),
            "saving": round(old - new, 2), "url": link, "clearance": False, "store": store, "category": cat,
        }
    return products


JWT_RE = re.compile(r"eyJ[A-Za-z0-9_-]{8,}\.eyJ[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}")
_PUBLIC_TOKEN = {}


def _jwt_info(tok):
    try:
        part = tok.split(".")[1]
        part += "=" * (-len(part) % 4)
        return json.loads(base64.urlsafe_b64decode(part))
    except Exception:
        return None


def discover_public_token(cfg):
    """Ψάχνει στη σελίδα του Public (και στα scripts της) το κλειδί πρόσβασης που δίνει σε κάθε επισκέπτη."""
    site = cfg["public_site"]
    try:
        page = requests.get(site + "/", headers=HEADERS, timeout=30).text
    except Exception as e:
        log(f"Public: δεν άνοιξε η αρχική σελίδα ({e})")
        return None
    texts = [page]
    for src in re.findall(r'<script[^>]+src=["\']([^"\']+)["\']', page)[:15]:
        try:
            texts.append(requests.get(urljoin(site + "/", src), headers=HEADERS, timeout=30).text)
        except Exception:
            pass
    for t in texts:
        for tok in JWT_RE.findall(t):
            info = _jwt_info(tok)
            if info and ("findbar" in json.dumps(info) or "search" in json.dumps(info)):
                return tok
    return None


def get_public_token(cfg):
    tok = (cfg.get("public_bearer") or "").strip()
    if tok.lower().startswith("bearer "):
        tok = tok[7:].strip()
    if not tok:
        tok = _PUBLIC_TOKEN.get("auto") or discover_public_token(cfg)
        if tok:
            _PUBLIC_TOKEN["auto"] = tok
            log("Public: βρέθηκε αυτόματα το κλειδί πρόσβασης.")
    if not tok:
        raise RuntimeError("Public: χρειάζεται κλειδί πρόσβασης. Βάλε το στο config.json στη γραμμή "
                           "public_bearer (δες το README.txt).")
    info = _jwt_info(tok)
    exp = info.get("exp") if info else None
    if isinstance(exp, (int, float)) and exp - time.time() < 7 * 86400:
        log("Public: το κλειδί πρόσβασης λήγει σύντομα - θα χρειαστεί νέο.")
    return tok


def _public_request(cfg, page, session_id):
    headers = dict(HEADERS)
    headers.update({"Content-Type": "application/json", "Accept": "application/json",
                    "Origin": cfg["public_site"], "Referer": cfg["public_site"] + "/",
                    "Authorization": "Bearer " + get_public_token(cfg)})
    payload = {"query": ".p",
               "filters": [{"field": "seller", "value": ["Public"]},
                           {"field": "availability", "value": ["Άμεσα Διαθέσιμο"]}],
               "initial": False, "page": page, "size": cfg["public_page_size"],
               "sort": "auto_discount_percent", "dir": "desc", "session_id": session_id}
    r = requests.post(cfg["public_api_url"], json=payload, headers=headers, timeout=45)
    if r.status_code == 401:
        _PUBLIC_TOKEN.clear()
        raise RuntimeError("Public HTTP 401: το κλειδί πρόσβασης δεν έγινε δεκτό ή έληξε. "
                           "Βάλε νέο στη γραμμή public_bearer του config.json.")
    if r.status_code != 200:
        raise RuntimeError(f"Public HTTP {r.status_code}: {r.text[:200]}")
    return r.json()


def scan_public(cfg, test=False):
    session_id = "".join(random.choice(string.ascii_letters + string.digits) for _ in range(32))
    products = {}
    for page in range(1, (1 if test else cfg["public_max_pages"]) + 1):
        data = _public_request(cfg, page, session_id)
        if test:
            with open(os.path.join(APP_DIR, "public_debug.json"), "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=1)
        found = parse_public(data, cfg["public_site"], cfg)
        fresh = {c: v for c, v in found.items() if c not in products}
        log(f"Public σελίδα {page}: {len(fresh)} προϊόντα με έκπτωση")
        if not fresh:
            if page == 1:
                raise RuntimeError("Public: δεν διάβασα κανένα προϊόν (άλλαξε η μορφή των δεδομένων). "
                                   "Τρέξε PriceWatch.exe --test-public")
            break
        products.update(fresh)
        lowest = min((v["old"] - v["new"]) / v["old"] * 100 for v in fresh.values())
        if lowest < cfg["public_stop_below_percent"]:
            break
        time.sleep(cfg["page_delay_sec"])
    return products


def run_test_public(cfg):
    session_id = "".join(random.choice(string.ascii_letters + string.digits) for _ in range(32))
    data = _public_request(cfg, 1, session_id)
    with open(os.path.join(APP_DIR, "public_debug.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    keys = list(data.keys())[:15] if isinstance(data, dict) else f"λίστα με {len(data)} στοιχεία"
    log(f"Public: απάντηση OK. Κλειδιά: {keys}")
    items = _find_items(data)
    log(f"Public: βρέθηκαν {len(items)} προϊόντα στη λίστα")
    if items:
        d = _src(items[0])
        flat = " | ".join(f"{k}={str(v)[:40]}" for k, v in d.items()
                          if isinstance(v, (str, int, float, bool)) or v is None)
        log("Public: πεδία πρώτου προϊόντος: " + flat[:1500])
    prods = parse_public(data, cfg["public_site"], cfg)
    log(f"Public: διαβάστηκαν {len(prods)} προϊόντα με έκπτωση")
    for p in list(prods.values())[:6]:
        log(f"  -{(p['old']-p['new'])/p['old']*100:.0f}% | {p['name'][:60]} | {p['old']:.2f} -> {p['new']:.2f} | {p['url']}")
    log("Αποθηκεύτηκε και το public_debug.json δίπλα στο PriceWatch.exe")


# =========================================================== KOTSOVOLOS
def kots_url(cfg, cat, page, size):
    params = (f"pageNumber={page}&pageSize={size}&catalogId={cfg['kotsovolos_catalog_id']}"
              f"&langId={cfg['kotsovolos_lang_id']}&orderBy={cfg['kotsovolos_order_by']}")
    return (f"{cfg['kotsovolos_site']}/api/ext/getProductsByCategory?params={quote(params, safe='')}"
            f"&catId={cat}&storeId={cfg['kotsovolos_store_id']}&isCPage=false")


def kots_get(cfg, url):
    headers = dict(HEADERS)
    headers.update({"Accept": "application/json, text/plain, */*", "Referer": cfg["kotsovolos_site"] + "/"})
    r = requests.get(url, headers=headers, timeout=45)
    if r.status_code != 200:
        raise RuntimeError(f"Kotsovolos HTTP {r.status_code}: {r.text[:150]}")
    return r.json()


def kots_ids(data):
    out = set()
    for it in _find_items(data, 2):
        c = _find(_src(it), PUBLIC_CODE, _text, depth=2)
        if c:
            out.add(c)
    return out


KOTS_CACHE = os.path.join(DATA_DIR, "kots_categories.json")


def kots_resolve_cat(cfg, path):
    """Βρίσκει τον αριθμό κατηγορίας (catId) από τη σελίδα της κατηγορίας (script __NEXT_DATA__)."""
    try:
        cache = json.load(open(KOTS_CACHE, encoding="utf-8"))
    except Exception:
        cache = {}
    if path in cache:
        return cache[path]
    url = path if path.startswith("http") else cfg["kotsovolos_site"] + path
    page = fetch(url)
    m = re.search(r'<script id="__NEXT_DATA__"[^>]*>(.*?)</script>', page, re.S)
    if not m:
        raise RuntimeError("δεν βρέθηκαν δεδομένα κατηγορίας στη σελίδα")
    pp = json.loads(m.group(1)).get("props", {}).get("pageProps", {})
    cid = pp.get("catalogIdPLP") or (pp.get("resultList") or {}).get("categoryId")
    if not cid:
        raise RuntimeError("δεν βρέθηκε catId")
    cache[path] = str(cid)
    try:
        json.dump(cache, open(KOTS_CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return str(cid)


def scan_kotsovolos(cfg, test=False):
    cfg = apply_selection(cfg, "kotsovolos")
    products = {}
    size = cfg["kotsovolos_page_size"]
    cats = [(c, c) for c in cfg["kotsovolos_categories"]]
    for path in cfg.get("kotsovolos_category_paths", []):
        try:
            cats.append((kots_resolve_cat(cfg, path), path))
        except Exception as e:
            log(f"Kotsovolos: δεν βρήκα κατηγορία για {path}: {e}")
        time.sleep(1)
    done, first = set(), True
    for cat, label in cats:
        if cat in done:
            continue
        done.add(cat)
        seen = set()
        before = len(products)
        for page in range(1, (1 if test else cfg["kotsovolos_max_pages"]) + 1):
            try:
                data = kots_get(cfg, kots_url(cfg, cat, page, size))
            except Exception as e:
                if page == 1 and size != 15:
                    log(f"Kotsovolos: πρόβλημα με pageSize={size} ({e}). Δοκιμάζω 15.")
                    size = 15
                    data = kots_get(cfg, kots_url(cfg, cat, page, size))
                else:
                    log(f"Kotsovolos {label} σελίδα {page}: {e}")
                    break
            ids = kots_ids(data)
            new_ids = ids - seen
            found = parse_public(data, cfg["kotsovolos_site"], cfg, store="Kotsovolos", prefix="kots:",
                                 depth=2, fallback=f"{cfg['kotsovolos_site']}/")
            nice = cfg.get("category_labels", {}).get(label) or (pretty_slug(label.rstrip("/").split("/")[-1]) if "/" in label else "")
            for fp in found.values():
                fp["category"] = nice
            if test:
                with open(os.path.join(APP_DIR, "kotsovolos_debug.json"), "w", encoding="utf-8") as f:
                    json.dump(data, f, ensure_ascii=False, indent=1)
                items = _find_items(data, 2)
                log(f"Kotsovolos δοκιμή: κατηγορία {cat} ({label}) -> {len(items)} προϊόντα στη σελίδα, {len(found)} με έκπτωση")
                if first and items:
                    d = _src(items[0])
                    flat = " | ".join(f"{k}={str(v)[:35]}" for k, v in d.items()
                                      if isinstance(v, (str, int, float, bool)) or v is None)
                    log("Kotsovolos: πεδία πρώτου προϊόντος: " + flat[:900])
                for p in list(found.values())[:3]:
                    log(f"  -{(p['old']-p['new'])/p['old']*100:.0f}% | {p['name'][:60]} | {p['old']:.2f} -> {p['new']:.2f} | {p['url']}")
                first = False
            if not new_ids:
                break
            seen |= ids
            products.update(found)
            time.sleep(cfg["page_delay_sec"])
        log(f"Kotsovolos κατηγορία {cat} ({label.split('/')[-1]}): {len(seen)} προϊόντα, {len(products) - before} με έκπτωση")
    if not products and not test:
        raise RuntimeError("Kotsovolos: δεν διάβασα κανένα προϊόν με έκπτωση. Τρέξε PriceWatch.exe --test-kotsovolos")
    return products


# ============================================== ΚΑΤΗΓΟΡΙΕΣ ΑΠΟ EXCEL
CATS_XLSX = os.path.join(APP_DIR, "kategories.xlsx")
CATS_CACHE = os.path.join(DATA_DIR, "kategories_cache.json")


def load_selected_categories():
    """Αν υπάρχει το kategories.xlsx δίπλα στο PriceWatch.exe, παίρνει τις γραμμές με 'ναι' στη στήλη D.
    Επιστρέφει (germanos_urls, kotsovolos_paths) ή None αν δεν υπάρχει αρχείο."""
    if not os.path.exists(CATS_XLSX):
        return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(CATS_XLSX, read_only=True, data_only=True)
        g_urls, k_paths, labels = [], [], {}
        for ws in wb.worksheets:
            sheet = ws.title.lower()
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(row) < 5:
                    continue
                flag = str(row[3] or "").strip().lower()
                if flag not in ("ναι", "nai", "yes", "x", "1", "✓"):
                    continue
                val = str(row[4] or "").strip().rstrip(",").strip().strip('"').strip()
                if not val:
                    continue
                labels[val] = " › ".join(x for x in (str(row[1] or "").strip(), str(row[2] or "").strip()) if x)
                if sheet.startswith("germ"):
                    g_urls.append(val)
                elif sheet.startswith("kots"):
                    k_paths.append(val)
        wb.close()
        try:
            json.dump({"germanos": g_urls, "kotsovolos": k_paths, "labels": labels},
                      open(CATS_CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        except Exception:
            pass
        return g_urls, k_paths, labels
    except Exception as e:
        log(f"Δεν μπόρεσα να διαβάσω το kategories.xlsx ({e}). Χρησιμοποιώ την τελευταία γνωστή λίστα.")
        try:
            d = json.load(open(CATS_CACHE, encoding="utf-8"))
            return d["germanos"], d["kotsovolos"], d.get("labels", {})
        except Exception:
            return None


def apply_selection(cfg, store):
    """Αντικαθιστά τις κατηγορίες του config με αυτές που έχεις επιλέξει στο kategories.xlsx."""
    sel = load_selected_categories()
    if sel is None:
        return cfg
    cfg = dict(cfg)
    g, k, labels = sel
    cfg["category_labels"] = dict(labels)
    if store == "germanos" and g:
        cfg["germanos_urls"] = [u if u.startswith("http") else "https://www.germanos.gr" + u for u in g]
        for u in g:
            if not u.startswith("http"):
                cfg["category_labels"]["https://www.germanos.gr" + u] = labels.get(u, "")
        log(f"Germanos: {len(g)} κατηγορίες από το kategories.xlsx")
    if store == "kotsovolos" and k:
        cfg["kotsovolos_category_paths"] = k
        cfg["kotsovolos_categories"] = []
        log(f"Kotsovolos: {len(k)} κατηγορίες από το kategories.xlsx")
    return cfg


# ============================================================ GERMANOS
PRICE_RE = re.compile(r"(\d{1,3}(?:\.\d{3})*,\d{2})\s*€")
AVAIL_RE = re.compile(r"\s(?:Μη Διαθέσιμο|Διαθέσιμο|Χαμηλή Διαθεσιμότητα|Εξαντλήθηκε|Προπαραγγελία)(?=\s|$)|\s\d{1,3}(?:\.\d{3})*,\d{2}\s*€")


def parse_germanos(page_html, page_url):
    """Επιστρέφει (προϊόντα_με_έκπτωση, όλα_τα_productId_της_σελίδας). Κάθε προϊόν είναι link με ?productId=..."""
    soup = BeautifulSoup(page_html, "html.parser")
    seen, found = set(), {}
    for a in soup.find_all("a", href=True):
        m = re.search(r"productId=(\d+)", a["href"])
        if not m:
            continue
        pid = m.group(1)
        text = clean_str(a.get_text(" ", strip=True))
        if pid in seen and pid in found:
            continue
        prices = PRICE_RE.findall(text)
        if not prices:
            seen.add(pid)
            continue
        seen.add(pid)
        if "Μη Διαθέσιμο" in text or len(prices) < 2:
            continue
        old, new = to_float(prices[0]), to_float(prices[1])
        if not (old > new > 0):
            continue
        cut = AVAIL_RE.search(text)
        head = (text[:cut.start()] if cut else text).strip()
        words = head.split()
        key = " ".join(words[:4])
        idx = head.find(key, len(key)) if key else -1
        name = head[idx:] if idx > 0 else head
        url = urljoin(page_url, a["href"]).split("#")[0]
        found[pid] = {
            "code": "germanos:" + pid, "name": name[:140], "old": old, "new": new,
            "saving": round(old - new, 2), "url": url, "clearance": False, "store": "Germanos", "category": "",
        }
    return {"germanos:" + k: v for k, v in found.items()}, seen


def germanos_page_url(cfg, base, n):
    url = base + cfg["germanos_sort"]
    return url + (f"&pagenum={n}" if n > 1 else "")


def scan_germanos(cfg, test=False):
    cfg = apply_selection(cfg, "germanos")
    products = {}
    for base in cfg["germanos_urls"]:
        all_ids = set()
        before = sum(1 for k in products if k.startswith("germanos:"))
        for n in range(1, (1 if test else cfg["germanos_max_pages"]) + 1):
            url = germanos_page_url(cfg, base, n)
            try:
                ph = fetch(url)
            except Exception as e:
                if "404" not in str(e):
                    log(f"Germanos {base.split('?')[0].rstrip('/').split('/')[-1]} σελίδα {n}: {e}")
                break
            found, ids = parse_germanos(ph, url)
            nice = cfg.get("category_labels", {}).get(base) or pretty_slug(base.split("?")[0].rstrip("/").split("/")[-1])
            for fp in found.values():
                fp["category"] = nice
            new_ids = ids - all_ids
            if test:
                log(f"Germanos δοκιμή: {base.split('?')[0].split('/')[-2]} -> {len(ids)} προϊόντα στη σελίδα, {len(found)} με έκπτωση")
                for p in list(found.values())[:4]:
                    log(f"  -{(p['old']-p['new'])/p['old']*100:.0f}% | {p['name'][:60]} | {p['old']:.2f} -> {p['new']:.2f} | {p['url']}")
            if not new_ids:
                break
            all_ids |= ids
            products.update(found)
            time.sleep(cfg["page_delay_sec"])
        log(f"Germanos {base.split('?')[0].rstrip('/').split('/')[-1]}: {len(all_ids)} προϊόντα, "
            f"{sum(1 for k in products if k.startswith('germanos:')) - before} με έκπτωση")
    if not products and not test:
        raise RuntimeError("Germanos: δεν διάβασα κανένα προϊόν με έκπτωση")
    return products


def scan_all(cfg):
    """Τα καταστήματα διαβάζονται ταυτόχρονα (το καθένα με τα δικά του διαλείμματα), άρα ο έλεγχος κρατάει όσο το πιο αργό."""
    jobs = [("Plaisio", scan)]
    if cfg.get("public_enabled"):
        jobs.append(("Public", scan_public))
    if cfg.get("kotsovolos_enabled"):
        jobs.append(("Kotsovolos", scan_kotsovolos))
    if cfg.get("germanos_enabled"):
        jobs.append(("Germanos", scan_germanos))
    products, errors = {}, []
    with ThreadPoolExecutor(max_workers=len(jobs)) as ex:
        futures = [(name, ex.submit(fn, cfg)) for name, fn in jobs]
        for name, fut in futures:
            try:
                products.update(fut.result())
            except Exception as e:
                errors.append(f"{name}: {e}")
                log(f"ΣΦΑΛΜΑ {name}: {e}")
    if not products:
        raise RuntimeError(" | ".join(errors) or "Κανένα προϊόν")
    return products


# ---------------------------------------------------------------- database
def db_connect():
    con = sqlite3.connect(DB_PATH)
    con.executescript("""
        CREATE TABLE IF NOT EXISTS prices(code TEXT, ts TEXT, price REAL, old_price REAL);
        CREATE INDEX IF NOT EXISTS ix_prices ON prices(code, ts);
        CREATE TABLE IF NOT EXISTS alerts(code TEXT, ts TEXT, price REAL);
        CREATE INDEX IF NOT EXISTS ix_alerts ON alerts(code, ts);
        CREATE TABLE IF NOT EXISTS state(k TEXT PRIMARY KEY, v TEXT);
    """)
    return con


def evaluate(con, cfg, products):
    now = datetime.now()
    since = (now - timedelta(days=30)).isoformat()
    recent = (now - timedelta(days=3)).isoformat()
    deals = []
    for p in products.values():
        if p["new"] < cfg["min_new_price_eur"]:
            continue
        disc = (p["old"] - p["new"]) / p["old"] * 100
        hist = [r[0] for r in con.execute("SELECT price FROM prices WHERE code=? AND ts>=?", (p["code"], since))]
        drop_hist = 0.0
        median = None
        if len(hist) >= 3:
            median = statistics.median(hist)
            drop_hist = (median - p["new"]) / median * 100 if median else 0
        hit_shop = disc >= cfg["min_discount_percent"]
        hit_hist = drop_hist >= cfg["min_drop_vs_history_percent"]
        if not (hit_shop or hit_hist):
            continue
        # να μην ειδοποιεί ξανά και ξανά για την ίδια τιμή
        last = con.execute("SELECT price FROM alerts WHERE code=? AND ts>=? ORDER BY ts DESC LIMIT 1",
                           (p["code"], recent)).fetchone()
        if last and p["new"] >= last[0] * 0.99:
            continue
        p.update(disc=disc, drop_hist=drop_hist, median=median, hit_shop=hit_shop, hit_hist=hit_hist)
        deals.append(p)
    deals.sort(key=lambda x: max(x["disc"], x["drop_hist"]), reverse=True)
    return deals


def record(con, products):
    ts = datetime.now().isoformat()
    con.executemany("INSERT INTO prices VALUES(?,?,?,?)",
                    [(p["code"], ts, p["new"], p["old"]) for p in products.values()])
    con.execute("DELETE FROM prices WHERE ts < ?", ((datetime.now() - timedelta(days=90)).isoformat(),))
    con.commit()


def format_deal(p):
    why = []
    if p["hit_shop"]:
        why.append(f"-{p['disc']:.0f}% από την τιμή του καταστήματος (€{p['old']:.2f})")
    if p["hit_hist"]:
        why.append(f"-{p['drop_hist']:.0f}% από τη συνήθη τιμή των τελευταίων ημερών (€{p['median']:.2f})")
    tag = " 🏷 CLEARANCE" if p["clearance"] else ""
    return (f"🔥 <b>{html.escape(p['name'])}</b>{tag}\n"
            f"💶 Τώρα: <b>€{p['new']:.2f}</b>\n"
            f"📉 " + " | ".join(why) + f"\n"
            f"🏪 {p.get('store', 'Plaisio')} · κωδ. {p['code'].split(':')[-1]}\n"
            f"🔗 {html.escape(p['url'])}")


def ranked(cfg, products, min_price):
    """Όλα τα προϊόντα με έκπτωση, ταξινομημένα κατά ΠΟΣΟΣΤΟ έκπτωσης (φθίνουσα)."""
    rows = []
    for p in products.values():
        if p["new"] < min_price:
            continue
        q = dict(p)
        q["pct"] = (p["old"] - p["new"]) / p["old"] * 100
        rows.append(q)
    rows.sort(key=lambda x: x["pct"], reverse=True)
    return rows


def tier_color(pct):
    """Χρώμα κελιού ανάλογα με το ποσοστό έκπτωσης (σε %)."""
    if pct >= 70:
        return "FF6666"   # κόκκινο
    if pct >= 50:
        return "FFC266"   # πορτοκαλί
    if pct >= 30:
        return "FFF2A8"   # κίτρινο
    return None


def _try_write(label, fn):
    """Αν ένα αρχείο είναι ανοιχτό στο Excel, τα Windows δεν επιτρέπουν εγγραφή - να μη σταματά το πρόγραμμα."""
    try:
        fn()
        return True
    except PermissionError:
        log(f"Δεν μπόρεσα να ενημερώσω το {label} (είναι ανοιχτό; κλείσ' το και θα γίνει στον επόμενο έλεγχο).")
    except Exception as e:
        log(f"Σφάλμα στο {label}: {e}")
    return False


def write_reports(cfg, products):
    rows = ranked(cfg, products, cfg["report_min_price_eur"])
    stamp = f"{datetime.now():%d/%m/%Y %H:%M}"
    fm = lambda v: f"{v:.2f}".replace(".", ",")
    store = lambda p: p.get("store", "Plaisio")
    cat = lambda p: clean_str(p.get("category", "") or "")
    shown = lambda p: p["code"].split(":")[-1]

    def write_csv():
        with open(os.path.join(DATA_DIR, "deals.csv"), "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Έκπτωση %", "Κατάστημα", "Κατηγορία", "Προϊόν", "Τιμή τώρα €", "Παλιά τιμή €", "Κέρδος €",
                        "Κωδικός", "Clearance", "Link"])
            for p in rows:
                w.writerow([f"{p['pct']:.0f}%", store(p), cat(p), p["name"], fm(p["new"]), fm(p["old"]),
                            fm(p["saving"]), shown(p), "ναι" if p["clearance"] else "", p["url"]])

    def write_html():
        body = "".join(
            f"<tr><td class=p{(' style=background:#' + tier_color(p['pct'])) if tier_color(p['pct']) else ''}>-{p['pct']:.0f}%</td>"
            f"<td>{store(p)}</td><td>{html.escape(cat(p))}</td>"
            f"<td><a href='{html.escape(p['url'])}' target=_blank>{html.escape(p['name'])}</a>"
            f"{' <b>CLEARANCE</b>' if p['clearance'] else ''}<br><small>Κωδ. {shown(p)}</small></td>"
            f"<td class=n>€{fm(p['new'])}</td><td class=n><s>€{fm(p['old'])}</s></td>"
            f"<td class=n>€{fm(p['saving'])}</td></tr>" for p in rows)
        page = ("<!doctype html><html lang=el><head><meta charset=utf-8>"
                "<meta name=viewport content='width=device-width,initial-scale=1'>"
                "<title>Προσφορές κατά % έκπτωσης</title><style>"
                "body{font-family:Segoe UI,Arial,sans-serif;margin:12px}table{border-collapse:collapse;width:100%}"
                "td,th{border-bottom:1px solid #ddd;padding:6px;text-align:left;vertical-align:top}"
                "th{background:#222;color:#fff;position:sticky;top:0}"
                ".p{font-weight:700;color:#c00;font-size:1.2em;white-space:nowrap}"
                ".n{white-space:nowrap}a{text-decoration:none;color:#06c}</style></head><body>"
                f"<h2>Προϊόντα ταξινομημένα κατά % έκπτωσης</h2><p>Ενημερώθηκε {stamp} · {len(rows)} προϊόντα "
                f"(τιμή ≥ €{cfg['report_min_price_eur']})</p><table><tr><th>%</th><th>Κατάστημα</th><th>Κατηγορία</th>"
                f"<th>Προϊόν</th><th>Τώρα</th><th>Παλιά</th><th>Κέρδος</th></tr>{body}</table></body></html>")
        with open(os.path.join(DATA_DIR, "deals.html"), "w", encoding="utf-8") as f:
            f.write(page)

    def write_xlsx():
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "Προσφορές"
        head = ["Έκπτωση %", "Κατάστημα", "Κατηγορία", "Προϊόν", "Τιμή τώρα €", "Παλιά τιμή €", "Κέρδος €",
                "Κωδικός", "Clearance", "Link"]
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="222222")
            c.alignment = Alignment(horizontal="center")
        for p in rows:
            ws.append([round(p["pct"] / 100, 4), store(p), cat(p), clean_str(p["name"]), p["new"], p["old"],
                       p["saving"], clean_str(shown(p)), "ναι" if p["clearance"] else "", clean_str(p["url"])])
        last = len(rows) + 1
        for r in range(2, last + 1):
            ws.cell(r, 1).number_format = "0%"
            ws.cell(r, 1).font = Font(bold=True, size=12)
            ws.cell(r, 1).alignment = Alignment(horizontal="center")
            tc = tier_color(rows[r - 2]["pct"])
            if tc:
                ws.cell(r, 1).fill = PatternFill("solid", fgColor=tc)
            for col in (5, 6, 7):
                ws.cell(r, col).number_format = '#,##0.00 "€"'
            ws.cell(r, 10).hyperlink = ws.cell(r, 10).value
            ws.cell(r, 10).font = Font(color="0563C1", underline="single")
        for i, w in enumerate([12, 12, 34, 70, 14, 14, 12, 14, 11, 60], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:J{max(last, 2)}"
        ws["L1"] = f"Ενημερώθηκε {stamp}"
        for rr, (txt, col) in enumerate([("70% και πάνω", "FF6666"), ("50% - 69%", "FFC266"), ("30% - 49%", "FFF2A8")], 2):
            ws.cell(rr, 12, txt).fill = PatternFill("solid", fgColor=col)
        ws.column_dimensions["L"].width = 24
        wb.save(os.path.join(DATA_DIR, "deals.xlsx"))

    _try_write("deals.html", write_html)
    _try_write("deals.csv", write_csv)
    _try_write("deals.xlsx", write_xlsx)
    log(f"Αναφορά ενημερώθηκε: deals.xlsx / deals.html ({len(rows)} προϊόντα, κατά % έκπτωσης)")


def maybe_digest(con, cfg, products):
    """Μία φορά τη μέρα στέλνει στο Telegram τα Top N κατά ποσοστό."""
    n = cfg["daily_digest_top_n"]
    if not n:
        return
    today = datetime.now().strftime("%Y-%m-%d")
    r = con.execute("SELECT v FROM state WHERE k='digest'").fetchone()
    if r and r[0] == today:
        return
    rows = ranked(cfg, products, cfg["min_new_price_eur"])[:n]
    if not rows:
        return
    lines = [f"{i}. <b>-{p['pct']:.0f}%</b> · €{p['new']:.2f} (από €{p['old']:.2f})\n"
             f"[{p.get('store', 'Plaisio')}] {html.escape(p['name'][:70])}\n{html.escape(p['url'])}" for i, p in enumerate(rows, 1)]
    if tg_send(cfg, f"📊 Top {len(rows)} εκπτώσεις κατά % (τιμή ≥ €{cfg['min_new_price_eur']})\n\n" + "\n\n".join(lines)):
        con.execute("INSERT OR REPLACE INTO state VALUES('digest', ?)", (today,))
        con.commit()
        if cfg.get("daily_xlsx_to_telegram"):
            tg_send_document(cfg, os.path.join(DATA_DIR, "deals.xlsx"), "Η πλήρης λίστα προσφορών (Excel)")


def run_once(cfg):
    log("Ξεκινά έλεγχος...")
    products = scan_all(cfg)
    con = db_connect()
    deals = evaluate(con, cfg, products)
    record(con, products)
    write_reports(cfg, products)
    limit = cfg["max_alerts_per_run"]
    for p in deals[:limit]:
        if tg_send(cfg, format_deal(p)):
            con.execute("INSERT INTO alerts VALUES(?,?,?)", (p["code"], datetime.now().isoformat(), p["new"]))
            con.commit()
        time.sleep(1)
    if len(deals) > limit:
        tg_send(cfg, f"… και άλλες {len(deals) - limit} προσφορές πάνω από το όριο.")
    maybe_digest(con, cfg, products)
    log(f"Τέλος: {len(products)} προϊόντα, {len(deals)} ειδοποιήσεις.")


def run_test(cfg):
    products = scan(cfg, test=True)
    log("--- ΔΟΚΙΜΗ: πρώτα 8 προϊόντα ---")
    for p in list(products.values())[:8]:
        log(f"{p['code']} | {p['name'][:60]} | παλιά {p['old']:.2f} -> νέα {p['new']:.2f} "
            f"(-{(p['old']-p['new'])/p['old']*100:.0f}%) | {p['url']}")
    best = sorted(products.values(), key=lambda x: (x['old'] - x['new']) / x['old'], reverse=True)[:5]
    write_reports(cfg, products)
    log("--- μεγαλύτερο % στη σελίδα 1 ---")
    for p in best:
        log(f"-{(p['old']-p['new'])/p['old']*100:.0f}% | {p['name'][:60]} | {p['new']:.2f}€")
    tpl = find_page_url_builder(cfg, open(os.path.join(APP_DIR, 'page1_debug.html'), encoding='utf-8').read(), cfg['start_url'])
    log(f"Σελιδοποίηση: {'βρέθηκε -> ' + tpl(2) if tpl else 'ΔΕΝ βρέθηκε (θα χρειαστεί page_url_template)'}")


def main():
    cfg = load_config()
    args = sys.argv[1:]
    if "--test" in args:
        run_test(cfg)
        return
    if "--test-public" in args:
        run_test_public(cfg)
        return
    if "--test-kotsovolos" in args:
        scan_kotsovolos(cfg, test=True)
        return
    if "--test-germanos" in args:
        scan_germanos(cfg, test=True)
        return
    detect_chat_id(cfg)
    if not cfg["telegram_token"] or not cfg["telegram_chat_id"]:
        log("Λείπει το Telegram token/chat id στο config.json. Δες το README.txt.")
        try:
            input("Πάτα Enter για έξοδο...")
        except EOFError:
            pass
        sys.exit(1)
    if "--once" in args:
        try:
            run_once(cfg)
        except Exception as e:
            log(f"ΣΦΑΛΜΑ: {e}")
            sys.exit(1)
        return
    tg_send(cfg, "✅ Το PriceWatch ξεκίνησε. Θα ελέγχει κάθε "
                 f"{cfg['check_every_minutes']} λεπτά για πτώσεις άνω του {cfg['min_discount_percent']}%.")
    failures = 0
    while True:
        started = time.time()
        try:
            run_once(cfg)
            failures = 0
        except Exception as e:
            failures += 1
            log(f"ΣΦΑΛΜΑ: {e}")
            if failures == 3:
                tg_send(cfg, f"⚠️ Το PriceWatch αποτυγχάνει 3 φορές στη σειρά: {e}")
        # ο χρόνος μετριέται από την ΑΡΧΗ του ελέγχου, όχι από το τέλος του
        took = time.time() - started
        wait = max(60, cfg["check_every_minutes"] * 60 - took)
        log(f"Ο έλεγχος κράτησε {took/60:.1f} λεπτά. Επόμενος έλεγχος σε {wait/60:.0f} λεπτά.")
        time.sleep(wait)


if __name__ == "__main__":
    main()
